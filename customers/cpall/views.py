"""
views.py — หน้าเว็บทั้งหมดของโมดูล cpall (CP All / 7-11)
"""
import io
import json
import os
import zipfile
from datetime import date, datetime
from urllib.parse import quote

from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse

from customers.cpall.forms import ImportPOForm, TemplateUploadForm
from customers.cpall.logic.excel_export import (
    find_po_barcodes_missing_in_template,
    find_po_sub_locations_missing_in_template,
)
from customers.cpall.logic.grouping import (
    DuplicateSubLocationError,
    InactiveSkuOrderedError,
    ReconciliationError,
    check_duplicate_sub_locations,
    check_inactive_skus_ordered,
)
from customers.cpall.logic.location_mapping_manager import get_existing_groups, save_location_mapping
from customers.cpall.logic.logistic_plan_export import (
    LogisticPlanError,
    get_group_templates,
    group_has_data,
    validate_logistic_plan,
)
from customers.cpall.logic.plan_regenerator import (
    PlanRegenerateError,
    regenerate_logistic_plan_bytes,
    regenerate_production_plan_bytes,
)
from customers.cpall.logic.plan_runner import (
    delete_plan_run,
    edit_buffer_and_regenerate,
    get_plan_run_detail,
    list_plan_runs,
    list_plan_runs_paginated,
    run_plan,
)
from customers.cpall.logic.plan_view_data import (
    get_logistic_plan_table,
    get_logistic_plan_table_from_db,
    get_production_plan_table,
    get_production_plan_table_from_db,
    get_skipped_skus,
)
from customers.cpall.logic.po_parser import (
    POInUseError,
    POParseError,
    check_duplicate_rows,
    check_unknown_locations,
    check_unknown_skus,
    delete_po_import,
    list_po_imports,
    list_po_imports_paginated,
    load_po_to_db,
)
from customers.cpall.logic.po_regenerator import PORegenerateError, regenerate_po_file_bytes
from customers.cpall.logic.po_view_data import get_po_detail
from customers.cpall.logic.product_master_manager import save_product
from customers.cpall.logic.template_manager import (
    TemplateInUseError,
    TemplateValidationError,
    activate_template_group,
    apply_product_master_action,
    build_group_reconcile_actions,
    delete_version,
    get_group_template_versions,
    get_template_grid,
    get_template_registry,
    list_templates,
    list_versions,
    reconcile_template_version,
    restore_to_version,
    upload_new_version,
    validate_group_consistency,
)
from customers.cpall.models import PlanRun, PoImport

UPLOAD_DIR = "customers/cpall/data/po_uploads"
TEMP_UPLOAD_DIR = "customers/cpall/data/temp_uploads"


def index(request):
    """แดชบอร์ด — สรุปภาพรวมสั้นๆ เท่านั้น (3 PO ล่าสุด + 3 แผนล่าสุด) ดูทั้งหมดแยกไปหน้า /po/ และ /plans/"""
    po_imports = list_po_imports(limit=3)
    plan_runs = list_plan_runs(limit=3)
    return render(request, "cpall/index.html", {"po_imports": po_imports, "plan_runs": plan_runs})


def po_list(request):
    """หน้ารายการ PO ทั้งหมด — นำเข้าใหม่/ติ๊กเลือกสร้างแผน/ลบ/ค้นหา/แบ่งหน้า ได้จากหน้านี้"""

    from urllib.parse import urlencode

    page = int(request.GET.get("page", 1) or 1)
    page_size = int(request.GET.get("page_size", 10) or 10)
    search = request.GET.get("q", "").strip()

    result = list_po_imports_paginated(page=page, page_size=page_size, search=search)

    base_qs = ("&" + urlencode({"q": search})) if search else ""

    import_warning = request.session.pop("import_warning", None)

    return render(request, "cpall/po_list.html", {
        "po_imports": result["items"],
        "total": result["total"],
        "page": result["page"],
        "page_size": result["page_size"],
        "total_pages": result["total_pages"],
        "search": search,
        "base_qs": base_qs,
        "import_warning": import_warning,
    })


def plan_list(request):
    """หน้ารายการแผนทั้งหมด — กด "สร้างแผนใหม่" แล้วไปเลือก PO ที่หน้า /po/"""
    from customers.cpall.logic.date_utils import parse_date_arg

    page = int(request.GET.get("page", 1) or 1)
    page_size = int(request.GET.get("page_size", 10) or 10)
    status = request.GET.get("status", "all")
    date_from_str = request.GET.get("date_from", "").strip()
    date_to_str = request.GET.get("date_to", "").strip()
    try:
        date_from = parse_date_arg(date_from_str) if date_from_str else None
    except ValueError:
        date_from = None
    try:
        date_to = parse_date_arg(date_to_str) if date_to_str else None
    except ValueError:
        date_to = None

    result = list_plan_runs_paginated(page=page, page_size=page_size, status=status,
                                       date_from=date_from, date_to=date_to)
    base_qs = f"&status={status}" if status != "all" else ""
    if date_from_str:
        base_qs += f"&date_from={date_from_str}"
    if date_to_str:
        base_qs += f"&date_to={date_to_str}"
    return render(request, "cpall/plan_list.html", {
        "plan_runs": result["items"], "total": result["total"], "page": result["page"],
        "page_size": result["page_size"], "total_pages": result["total_pages"],
        "status": status, "base_qs": base_qs,
        "date_from": date_from_str, "date_to": date_to_str,
    })


def import_form(request):
    return render(request, "cpall/import.html", {"form": ImportPOForm()})


def _post_import_destination(request, po_import_id):
    """
    หา URL ปลายทางหลัง import PO เสร็จ (Add PO feature — 2025-09-12)

    ปกติจบที่หน้า PO List เหมือนเดิมทุกประการ — ยกเว้นกรณีที่ผู้ใช้เข้ามาจาก flow "เพิ่ม PO เข้าแผน"
    (หน้า add_po_form ส่ง add_to_plan มาด้วย แล้วเก็บไว้ใน session) จะพากลับไปที่แผนนั้นต่อทันที
    พร้อม PO ที่เพิ่ง import — ไม่ให้หลุด context ระหว่างทาง (flow import อาจแวะหน้า resolve-locations
    ก่อนได้ เลยต้องใช้ session ไม่ใช่ query param เพื่อให้ context รอดข้ามหน้า)

    pop ทิ้งทุกครั้งที่ใช้ — ใช้ครั้งเดียวจบ ไม่ค้างไปรบกวน import รอบถัดไป
    """
    plan_id = request.session.pop("add_to_plan", None)
    if plan_id:
        return reverse("cpall:add_po_buffer", args=[plan_id]) + f"?new_po_ids={po_import_id}"
    return reverse("cpall:po_list")


def import_submit(request):
    if request.method != "POST":
        return redirect("cpall:import_form")

    is_htmx = request.headers.get("HX-Request") == "true"

    # มาจาก flow "เพิ่ม PO เข้าแผน" หรือเปล่า — เก็บไว้ใน session เพื่อให้รอดข้ามหน้า resolve-locations
    # ที่ import อาจแวะก่อน (Add PO feature — 2025-09-12) ไม่กระทบ flow import ปกติเลย ถ้าไม่ได้ส่งมา
    add_to_plan = request.POST.get("add_to_plan")
    if add_to_plan:
        try:
            request.session["add_to_plan"] = int(add_to_plan)
        except ValueError:
            pass

    form = ImportPOForm(request.POST, request.FILES)
    if not form.is_valid():
        # Django Form validation ไม่ผ่าน (กรอกไม่ครบ/ไฟล์ผิดชนิด) -> อยู่หน้าเดิม โชว์ error ในฟอร์มเลย
        # ผ่าน HTMX swap แค่ตัวฟอร์ม (_import_form.html) ไม่ reload ทั้งหน้า
        if is_htmx:
            return render(request, "cpall/_import_form.html", {"form": form})
        return render(request, "cpall/import.html", {"form": form})

    po_file = form.cleaned_data["po_file"]
    production_date = form.cleaned_data["production_date"]
    po_date = form.cleaned_data["po_date"]

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    original_name = os.path.basename(po_file.name)  # กันเผื่อชื่อไฟล์มี path แปลกปน
    saved_path = os.path.join(UPLOAD_DIR, f"{timestamp}_{original_name}")
    with open(saved_path, "wb") as f:
        for chunk in po_file.chunks():
            f.write(chunk)

    def error_response(message, status=400):
        # นำเข้าไม่สำเร็จ -> ไม่เก็บไฟล์ที่เพิ่งบันทึกไว้ค้างบนดิสก์ (กันพื้นที่บวมจากไฟล์ที่ import พลาด)
        if os.path.exists(saved_path):
            os.remove(saved_path)
        if is_htmx:
            response = HttpResponse(status=status)
            response["HX-Trigger"] = json.dumps({"toast": {"message": message, "level": "error"}})
            return response
        return render(request, "cpall/import_result.html",
                       {"success": False, "error": message, "original_filename": po_file.name})

    try:
        duplicate_groups = check_duplicate_rows(saved_path)
    except POParseError as e:
        return error_response(f"ไฟล์ PO มีปัญหา: {e}")
    except Exception as e:
        # ไฟล์ผิดประเภท (เช่น ไม่ใช่ .xlsx จริง) เป็นความผิดพลาดของผู้ใช้ (เลือกไฟล์ผิด) ไม่ใช่ระบบ
        # พังเอง — ใช้ 400 แทน 500 เพื่อให้ log/monitoring แยกแยะได้ถูกต้อง (เจอจากการทดสอบ 2025-09-06)
        return error_response(f"อ่านไฟล์ล้มเหลว: {type(e).__name__}: {e}", status=400)

    if duplicate_groups:
        # เจอรายการที่อาจซ้ำ (po_number+fc_code+barcode+line_no ตรงกันเป๊ะ) — ไม่ import ต่อทันที
        # ให้ Admin เห็นรายละเอียดแล้วเลือกเองว่า Continue (import ทุกแถวตามไฟล์จริง ไม่ตัดอะไรออก)
        # หรือ Stop (ยกเลิก ลบไฟล์ที่ค้างทิ้ง) — เก็บ path+วันที่ไว้ใน session (ไม่ใช่ hidden field ใน
        # form กัน path ถูกแก้ผ่าน browser dev tools ได้) รอ Admin ตัดสินใจที่หน้า confirm_duplicates
        request.session["pending_po_import"] = {
            "saved_path": saved_path,
            "production_date": production_date.isoformat(),
            "po_date": po_date.isoformat(),
            "original_filename": po_file.name,
        }
        context = {"duplicate_groups": duplicate_groups, "original_filename": po_file.name}
        if is_htmx:
            response = HttpResponse(status=200)
            response["HX-Redirect"] = reverse("cpall:confirm_duplicates")
            return response
        return render(request, "cpall/confirm_duplicates.html", context)

    try:
        po_import_id = load_po_to_db(saved_path, production_date, po_date, imported_by="web")
    except POParseError as e:
        return error_response(f"ไฟล์ PO มีปัญหา: {e}")
    except Exception as e:
        return error_response(f"นำเข้าล้มเหลว: {type(e).__name__}: {e}", status=500)

    # นำเข้าสำเร็จ -> ข้อมูลทุกคอลัมน์ของไฟล์ (ไม่ใช่แค่ 12 คอลัมน์ที่ระบบใช้คำนวณ) ถูกเก็บไว้ครบใน
    # po_import.column_order + po_line.all_values แล้ว (ดู po_parser.py's parse_po_file) ไม่จำเป็นต้อง
    # เก็บไฟล์ต้นฉบับไว้บนดิสก์อีกต่อไปเลย — ดาวน์โหลดย้อนหลังทีหลัง จะสร้างไฟล์ใหม่จากข้อมูลใน DB แทน
    # (ดู po_regenerator.py) ทดสอบยืนยันแล้วว่าตรงกับต้นฉบับทุกเซลล์ 100%
    if os.path.exists(saved_path):
        os.remove(saved_path)

    # ต่อ flow เดิม — เช็ค location ที่ไม่รู้จักก่อน
    unknown_locations = check_unknown_locations(po_import_id)
    if unknown_locations:
        if is_htmx:
            response = HttpResponse(status=200)
            response["HX-Redirect"] = reverse(
                "cpall:resolve_locations",
                args=[po_import_id],
            )
            return response
        return redirect(
            "cpall:resolve_locations",
            po_import_id=po_import_id,
        )

    # Product ตรวจจาก Active Production Template
    # ไม่ตรวจ ProductMaster และไม่พาไป resolve_products
    missing_in_template = find_po_barcodes_missing_in_template([po_import_id])

    if missing_in_template:
        detail_url = (
            reverse("cpall:missing_template_items")
            + f"?po_import_ids={po_import_id}"
        )

        request.session["import_warning"] = {
            "message": (
                f"นำเข้า PO สำเร็จ — พบสินค้า {len(missing_in_template)} รายการ "
                "ที่ไม่มีใน Template แพลนผลิต"
            ),
            "detail_url": detail_url,
        }


    destination = _post_import_destination(request, po_import_id)
    if is_htmx:
        response = HttpResponse(status=200)
        response["HX-Redirect"] = destination
        return response

    return redirect(destination)


def confirm_duplicates(request):
    """
    หน้า "พบรายการที่อาจซ้ำ" — Admin เลือก Continue (import ทุกแถวตามไฟล์จริง ไม่ตัดอะไรออกเลย)
    หรือ Stop (ยกเลิก ลบไฟล์ที่ค้างทิ้ง) — path/วันที่ที่ต้องใช้เก็บไว้ใน session จาก import_submit()
    ไม่รับผ่าน POST body เอง (กัน path ถูกแก้ผ่าน browser dev tools ได้)
    """
    pending = request.session.get("pending_po_import")
    if pending is None:
        return render(request, "cpall/plan_error.html",
                       {"error": "ไม่พบไฟล์ที่รออยู่ (session อาจหมดอายุ) กรุณาอัปโหลดใหม่"})

    is_htmx = request.headers.get("HX-Request") == "true"

    if request.method == "POST":
        action = request.POST.get("action")
        del request.session["pending_po_import"]  # ใช้ครั้งเดียวจบ ไม่ว่าจะกดปุ่มไหน

        if action == "stop":
            if os.path.exists(pending["saved_path"]):
                os.remove(pending["saved_path"])
            if is_htmx:
                response = HttpResponse(status=200)
                response["HX-Trigger"] = json.dumps(
                    {"toast": {"message": "ยกเลิกการนำเข้าแล้ว", "level": "success"}})
                response["HX-Redirect"] = reverse("cpall:import_form")
                return response
            return redirect("cpall:import_form")

        # action == "continue" (หรือค่าอื่นที่ไม่ใช่ stop ก็ถือว่า continue ไปเลย ปลอดภัยกว่าปฏิเสธ
        # การนำเข้าไปเฉยๆ โดยไม่มีเหตุผล) — import ทุกแถวจริง ไม่ตัดอะไรออกแม้จะซ้ำก็ตาม
        production_date = date.fromisoformat(pending["production_date"])
        po_date = date.fromisoformat(pending["po_date"])
        try:
            po_import_id = load_po_to_db(pending["saved_path"], production_date, po_date, imported_by="web")
        except POParseError as e:
            return render(request, "cpall/plan_error.html", {"error": f"ไฟล์ PO มีปัญหา: {e}"})
        except Exception as e:
            return render(request, "cpall/plan_error.html",
                           {"error": f"นำเข้าล้มเหลว: {type(e).__name__}: {e}"})

        if os.path.exists(pending["saved_path"]):
            os.remove(pending["saved_path"])

        # ต่อ flow เดิมเป๊ะ (เหมือน import_submit ตอนไม่มี duplicate) — เช็ค location/SKU ที่ไม่รู้จัก
        unknown_locations = check_unknown_locations(po_import_id)
        if unknown_locations:
            if is_htmx:
                response = HttpResponse(status=200)
                response["HX-Redirect"] = reverse("cpall:resolve_locations", args=[po_import_id])
                return response
            return redirect("cpall:resolve_locations", po_import_id=po_import_id)

        missing_in_template = find_po_barcodes_missing_in_template([po_import_id])

        if missing_in_template:
            detail_url = (
                reverse("cpall:missing_template_items")
                + f"?po_import_ids={po_import_id}"
            )

            request.session["import_warning"] = {
                "message": (
                    f"นำเข้า PO สำเร็จ — พบสินค้า {len(missing_in_template)} รายการ "
                    "ที่ไม่มีใน Template แพลนผลิต"
                ),
                "detail_url": detail_url,
            }
        # ไม่ใช่แล้ว ใช้ template จัดการเอา
        # unknown_skus = check_unknown_skus(po_import_id)
        # if unknown_skus:
        #     if is_htmx:
        #         response = HttpResponse(status=200)
        #         response["HX-Redirect"] = reverse("cpall:resolve_products", args=[po_import_id])
        #         return response
        #     return redirect("cpall:resolve_products", po_import_id=po_import_id)

        destination = _post_import_destination(request, po_import_id)
        if is_htmx:
            response = HttpResponse(status=200)
            response["HX-Redirect"] = destination
            return response
        return redirect(destination)

    # GET — แสดงหน้ายืนยัน (เรียก check_duplicate_rows ใหม่อีกครั้งจากไฟล์ที่ยังค้างอยู่ เผื่อ Admin
    # รีเฟรชหน้านี้ — ไม่ query จาก session เพราะ session เก็บแค่ path ไม่ได้เก็บรายละเอียดกลุ่มที่ซ้ำ)
    try:
        duplicate_groups = check_duplicate_rows(pending["saved_path"])
    except Exception:
        duplicate_groups = []
    return render(request, "cpall/confirm_duplicates.html", {
        "duplicate_groups": duplicate_groups, "original_filename": pending["original_filename"],
    })


def resolve_locations(request, po_import_id):
    unknown_locations = check_unknown_locations(po_import_id)
    is_htmx = request.headers.get("HX-Request") == "true"

    if request.method == "POST":
        for fc_code, name_th in unknown_locations:
            group = request.POST.get(f"group_{fc_code}", "").strip()
            new_group = request.POST.get(f"new_group_{fc_code}", "").strip()
            sub_location = request.POST.get(f"sub_location_{fc_code}", "").strip()

            final_group = new_group if new_group else group
            if final_group and sub_location:
                save_location_mapping(fc_code, name_th, final_group, sub_location)

        remaining = check_unknown_locations(po_import_id)
        if remaining:
            # บางรายการยังไม่ได้กรอก (เว้นว่างไว้) -> โชว์ฟอร์มเดิมต่อ เฉพาะรายการที่เหลือ
            context = {
                "po_import_id": po_import_id, "unknown_locations": remaining,
                "existing_groups": get_existing_groups(),
                "error": "ยังมีบางรายการที่ยังไม่ได้เลือกกลุ่มพื้นที่/กรอกจุดส่งย่อย",
            }
            if is_htmx:
                return render(request, "cpall/_resolve_locations_form.html", context)
            return render(request, "cpall/resolve_locations.html", context)

        # ไม่ใช้แล้ว ใช้ template จัดการเอา
        # location ครบหมดแล้ว -> เช็คต่อว่ามีสินค้าที่ยังไม่รู้จักไหม (ไม่บังคับ แค่แนะนำ)
        # unknown_skus = check_unknown_skus(po_import_id)
        # if unknown_skus:
        #     if is_htmx:
        #         response = HttpResponse(status=200)
        #         response["HX-Redirect"] = reverse("cpall:resolve_products", args=[po_import_id])
        #         return response
        #     return redirect("cpall:resolve_products", po_import_id=po_import_id)

        destination = _post_import_destination(request, po_import_id)
        if is_htmx:
            response = HttpResponse(status=200)
            response["HX-Redirect"] = destination
            return response
        return redirect(destination)

    return render(request, "cpall/resolve_locations.html", {
        "po_import_id": po_import_id, "unknown_locations": unknown_locations,
        "existing_groups": get_existing_groups(),
    })


def resolve_products(request, po_import_id):
    """ให้ Admin จัดการสินค้าที่ยังไม่รู้จักให้ครบทุก SKU ก่อนจึงไปต่อได้"""
    unknown_skus = check_unknown_skus(po_import_id)
    is_htmx = request.headers.get("HX-Request") == "true"

    if request.method == "POST":
        errors = []

        for barcode, item_name, net_case_price in unknown_skus:
            name_th = request.POST.get(f"name_th_{barcode}", "").strip() or item_name
            name_en = request.POST.get(f"name_en_{barcode}", "").strip()
            pack_size_str = request.POST.get(f"pack_size_{barcode}", "").strip()
            unit_price_str = request.POST.get(f"unit_price_{barcode}", "").strip()

            # ต้องกรอก Pack Size
            if not pack_size_str:
                errors.append(f"{item_name} ({barcode}): กรุณากรอก Pack Size")
                continue

            try:
                pack_size = int(pack_size_str)
                if pack_size <= 0:
                    raise ValueError

                # ราคาไม่ได้กรอกเอง → ใช้ราคาจาก PO
                unit_price = (
                    float(unit_price_str)
                    if unit_price_str
                    else net_case_price
                )

                if unit_price is None:
                    raise ValueError

            except ValueError:
                errors.append(
                    f"{item_name} ({barcode}): Pack Size หรือราคาสินค้าไม่ถูกต้อง"
                )
                continue

            save_product(
                barcode,
                name_th,
                name_en or None,
                pack_size,
                unit_price,
            )

        # ตรวจซ้ำหลังจากบันทึก
        remaining = check_unknown_skus(po_import_id)

        if remaining:
            context = {
                "po_import_id": po_import_id,
                "unknown_skus": remaining,
                "errors": errors,
            }

            if is_htmx:
                return render(
                    request,
                    "cpall/_resolve_products_form.html",
                    context,
                )

            return render(
                request,
                "cpall/resolve_products.html",
                context,
            )

        if is_htmx:
            response = HttpResponse(status=200)
            response["HX-Redirect"] = reverse("cpall:po_list")
            return response

        return redirect("cpall:po_list")

    return render(
        request,
        "cpall/resolve_products.html",
        {
            "po_import_id": po_import_id,
            "unknown_skus": unknown_skus,
        },
    )


def new_plan_submit(request):
    if request.method != "POST":
        return redirect("cpall:index")

    is_htmx = request.headers.get("HX-Request") == "true"

    def error_response(message, status=400, detail_url=None):
        if is_htmx:
            toast = {
                "message": message,
                "level": "error",
            }

            if detail_url:
                toast["detail_url"] = detail_url

            response = HttpResponse(status=status)
            response["HX-Trigger"] = json.dumps({"toast": toast})
            return response

        return render(request, "cpall/plan_error.html", {"error": message})

    po_import_ids = [int(x) for x in request.POST.getlist("po_import_ids")]
    if not po_import_ids:
        return error_response("ต้องเลือก PO อย่างน้อย 1 รอบ")
    
    # เช็ค Barcode ใน PO ว่ามีอยู่ใน Production Template หรือไม่
    missing_in_template = find_po_barcodes_missing_in_template(po_import_ids)

    if missing_in_template:
        names = ", ".join(
            f"{item['barcode']} ({item['product_name']})"
            for item in missing_in_template
        )

        detail_url = (
            reverse("cpall:missing_template_items")
            + "?"
            + "&".join(
                f"po_import_ids={po_import_id}"
                for po_import_id in po_import_ids
            )
        )

        return error_response(
            f"ไม่สามารถสร้างแผนได้ — พบ Barcode ใน PO "
            f"ที่ไม่มีใน Template แพลนผลิต จำนวน {len(missing_in_template)} รายการ"
            # f"\n{names}"
            f"\nกรุณาตรวจสอบและแก้ไข Template แพลนผลิต ก่อนสร้าง Plan",
            status=409,
            detail_url=detail_url,
        )

    # เช็ค Location ใน PO ว่ามีอยู่ใน Production Template หรือไม่
    missing_sub_locations = find_po_sub_locations_missing_in_template(
        po_import_ids
    )

    if missing_sub_locations:
        locations = ", ".join(missing_sub_locations)

        return error_response(
            f"ไม่สามารถสร้างแผนได้ — พบจุดส่งย่อยใน PO "
            f"ที่ไม่มีใน Template แพลนผลิต จำนวน "
            f"{len(missing_sub_locations)} รายการ: "
            f"{locations}"
            f"\nกรุณาตรวจสอบและแก้ไข Template แพลนผลิต ก่อนสร้าง Plan",
            status=409,
        )

    # เช็ค duplicate sub_location
    duplicate_sub_locations = check_duplicate_sub_locations(po_import_ids)

    if duplicate_sub_locations:
        details = "; ".join(
            f"{item['sub_location']} "
            f"(PO Import {', '.join(map(str, item['po_import_ids']))})"
            for item in duplicate_sub_locations
        )

        return error_response(
            f"ไม่สามารถสร้างแผนได้ — จุดส่งย่อยซ้ำกันในหลายรอบ PO: {details}",
            status=409,
        )

    # เช็ค Logistic Template เฉพาะ Group ที่มีข้อมูลใน PO ชุดนี้
    logistic_errors = []

    for group_name in get_group_templates():
        if not group_has_data(po_import_ids, group_name):
            continue

        try:
            result = validate_logistic_plan(
                po_import_ids,
                group_name,
            )
        except LogisticPlanError as e:
            logistic_errors.append(
                f"{group_name}: {e}"
            )
            continue

        if result["missing_sub_locations"]:
            locations = ", ".join(
                result["missing_sub_locations"]
            )

            logistic_errors.append(
                f"{group_name}: "
                f"ไม่มีจุดส่งย่อยใน Template "
                f"({locations})"
            )

        if result["overflow"]:
            details = "; ".join(
                f"{item['sub_location']} "
                f"(ต้องใช้ {item['needed']} PO / "
                f"รองรับ {item['available']} คอลัมน์)"
                for item in result["overflow"]
            )

            logistic_errors.append(
                f"{group_name}: "
                f"จำนวนคอลัมน์ PO ไม่พอ — {details}"
            )

        if result["missing_barcodes"]:
            logistic_errors.append(
                f"{group_name}: "
                f"พบ Barcode ที่ไม่มีใน Logistic Template "
                f"จำนวน {len(result['missing_barcodes'])} รายการ"
            )

    if logistic_errors:
        details = "\n".join(
            f"- {error}"
            for error in logistic_errors
        )

        return error_response(
            "ไม่สามารถสร้างแผนได้ — "
            "พบปัญหาใน Logistic Template:\n"
            f"{details}\n"
            "กรุณาตรวจสอบและแก้ไข Template ก่อนสร้าง Plan",
            status=409,
        )


    # เช็ค SKU ใหม่ที่ยังไม่มีใน ProductMaster
    unknown_skus = []
    for po_import_id in po_import_ids:
        unknown_skus.extend(check_unknown_skus(po_import_id))

    if unknown_skus:
        names = ", ".join(
            f"{barcode} ({item_name})"
            for barcode, item_name, _ in unknown_skus
        )
        return error_response(
            f"สร้างแผนไม่ได้ — พบสินค้าใหม่ที่ยังไม่ได้เพิ่มใน Product Master: {names}"
            # f"กรุณาจัดการสินค้าเหล่านี้ก่อน"
            f"\nตรวจสอบ Template ว่ามีสินค้าเหล่านี้อยู่หรือไม่ (ถ้าไม่มี ให้เพิ่มใน Template ก่อน)",
            status=409,
        )

    # เช็ค inactive SKU ที่ยังมี PO สั่งอยู่จริง
    inactive_ordered = check_inactive_skus_ordered(po_import_ids)

    # po_import_ids = [int(x) for x in request.POST.getlist("po_import_ids")]
    # if not po_import_ids:
    #     return error_response("ต้องเลือก PO อย่างน้อย 1 รอบ")

    # # เช็ค inactive SKU ที่ยังมี PO สั่งอยู่จริง "ก่อน" ไปหน้ากรอกยอดเผื่อเลย — เดิมเช็คแค่ตอน submit
    # # ยอดเผื่อ (ใน run_plan()) ทำให้ Admin ต้องกรอกยอดเผื่อครบ 19 SKU ก่อน ถึงจะรู้ว่าสร้างแผนไม่ได้
    # # เสียเวลาโดยไม่จำเป็น — ย้ายมาเช็คตรงนี้เพื่อบล็อกให้เร็วที่สุด (ยังคงเช็คซ้ำใน run_plan() ไว้ด้วย
    # # เผื่อกรณี SKU เพิ่งถูกปิดใช้งานระหว่างที่ Admin เปิดหน้ากรอกยอดเผื่อค้างไว้อยู่)
    # inactive_ordered = check_inactive_skus_ordered(po_import_ids)
    if inactive_ordered:
        names = ", ".join(f"{s['barcode']} ({s['name_th']})" for s in inactive_ordered)
        return error_response(
            f"สร้างแผนไม่สำเร็จ — PO รอบนี้สั่งสินค้าที่ถูกปิดใช้งานอยู่: {names}"
            # f"ไปเปิดใช้งาน (is_active) สินค้านี้ก่อนใน Django Admin ถึงจะสร้างแผนได้"
            f"\nตรวจสอบ Template ว่ามีสินค้าเหล่านี้อยู่หรือไม่ (ถ้าไม่มี ให้เพิ่มใน Template ก่อน)", status=409,
        )

    # ไปหน้ากรอกยอดเผื่อเสมอ — ขึ้นทุกครั้งที่สร้างแผน ไม่ใช่แค่ตอนมีรอบเช้าต่างจังหวัด
    # (ยอดเผื่อเป็นข้อมูลสำคัญที่ Admin ต้องยืนยันทุกรอบ ไม่ใช่แค่กลุ่มเช้าต่างจังหวัด)
    if is_htmx:
        response = HttpResponse(status=200)
        response["HX-Redirect"] = _buffer_form_url(po_import_ids)
        return response
    return redirect_to_buffer_form(request, po_import_ids)

def missing_template_items(request):
    po_import_ids = [
        int(x) for x in request.GET.getlist("po_import_ids")
    ]

    if not po_import_ids:
        return render(
            request,
            "cpall/missing_template_items.html",
            {"items": []},
        )

    items = find_po_barcodes_missing_in_template(po_import_ids)

    return render(
        request,
        "cpall/missing_template_items.html",
        {
            "items": items,
        },
    )

def _buffer_form_url(po_import_ids):
    from urllib.parse import urlencode
    qs = urlencode({"po_import_ids": ",".join(str(x) for x in po_import_ids)})
    return f"{reverse('cpall:buffer_form')}?{qs}"


def redirect_to_buffer_form(request, po_import_ids):
    return redirect(_buffer_form_url(po_import_ids))


def buffer_form(request):
    import openpyxl

    from customers.cpall.logic.excel_export import TEMPLATE_PATH as PP_TEMPLATE_PATH
    from customers.cpall.logic.excel_export import (
        _find_sku_header_rows as _find_pp_sku_header_rows,
    )
    from customers.cpall.logic.excel_export import (
        get_sheet_name as get_pp_sheet_name,
    )
    from customers.cpall.models import PlanSkuResult, ProductMaster

    po_import_ids_str = request.GET.get("po_import_ids", "")
    po_import_ids = [int(x) for x in po_import_ids_str.split(",") if x]
    if not po_import_ids:
        return render(request, "cpall/plan_error.html", {"error": "ไม่พบรอบ PO ที่เลือกไว้"})

    # ดึง SKU จาก ProductMaster ที่ยัง active อยู่เป็นหลัก (Database = source of truth) — เดิมดึงจาก
    # Production Plan template ตรงๆ ทำให้เพิ่ม SKU ใหม่ผ่าน Django Admin แล้วไม่ขึ้นในฟอร์มนี้เลย (ต้อง
    # ไปเพิ่มแถวในไฟล์ template ตรงๆ ถึงจะขึ้น) กลับหัวกลับหางกับที่ควรเป็น — แก้ให้ ProductMaster (DB)
    # เป็นตัวตัดสินว่า "มี SKU อะไรบ้าง" ส่วน Template ใช้แค่หาว่าแถวไหนอยู่ตรงไหน (คนละหน้าที่กัน) —
    # SKU ที่ active ใน DB แต่ไม่มีแถวใน Template เลย ไม่แสดงในฟอร์ม (ไม่มีที่เก็บค่าจริง) แต่เก็บไว้
    # แจ้งเตือน Admin ว่าขาดอะไรไป (2025-09-05)
    wb = openpyxl.load_workbook(PP_TEMPLATE_PATH)
    ws = wb[get_pp_sheet_name()]
    header_rows = _find_pp_sku_header_rows(ws)  # {barcode: row} — เอาไว้เรียงลำดับ + เช็คว่ามีแถวจริง

    active_barcodes = set(ProductMaster.objects.filter(is_active=True).values_list("barcode", flat=True))
    barcodes_in_template = set(header_rows.keys())
    # เรียงตามตำแหน่งแถวในไฟล์ (ตามลำดับใน Production Plan จริง) — เอาเฉพาะที่ active ใน DB ด้วย
    barcodes = [
        bc for bc, _ in sorted(header_rows.items(), key=lambda kv: kv[1])
        if bc in active_barcodes
    ]
    missing_in_template = active_barcodes - barcodes_in_template  # active ใน DB แต่ไม่มีแถวจริง

    # default ยอดเผื่อ: ดึงจาก buffer_qty ล่าสุดที่เคยบันทึกไว้ใน plan_sku_result —
    # ไม่ใช่จากไฟล์เทมเพลต (ซึ่งเป็นค่าเก่าที่ Admin กรอกไว้ครั้งแรก ไม่ใช่ล่าสุด)
    default_buffer = {}
    last_plan = PlanSkuResult.objects.filter(
        sheet_type="production", buffer_qty__isnull=False
    ).order_by("-id").first()
    if last_plan:
        for row in PlanSkuResult.objects.filter(
            plan_run_id=last_plan.plan_run_id, sheet_type="production", buffer_qty__isnull=False
        ).values("barcode", "buffer_qty").distinct("barcode"):
            if row["buffer_qty"] is not None:
                default_buffer[row["barcode"]] = float(row["buffer_qty"])

    name_lookup = {s.barcode: s.name_th for s in ProductMaster.objects.filter(barcode__in=barcodes)}
    sku_rows = [
        {"barcode": bc, "name_th": name_lookup.get(bc, bc), "default_buffer": default_buffer.get(bc, 0)}
        for bc in barcodes
    ]
    missing_names = list(
        ProductMaster.objects.filter(barcode__in=missing_in_template).values_list("name_th", flat=True)
    )

    return render(request, "cpall/buffer_form.html", {
        "po_import_ids": po_import_ids,
        "po_import_ids_str": po_import_ids_str,
        "sku_rows": sku_rows,
        "missing_names": missing_names,
    })


def buffer_form_submit(request):
    if request.method != "POST":
        return redirect("cpall:po_list")

    is_htmx = request.headers.get("HX-Request") == "true"

    def error_response(message, status=400):
        if is_htmx:
            response = HttpResponse(status=status)
            response["HX-Trigger"] = json.dumps({"toast": {"message": message, "level": "error"}})
            return response
        return render(request, "cpall/plan_error.html", {"error": message})

    po_import_ids_str = request.POST.get("po_import_ids_str", "")
    po_import_ids = [int(x) for x in po_import_ids_str.split(",") if x]
    if not po_import_ids:
        return error_response("ไม่พบรอบ PO ที่เลือกไว้")

    buffer_override = {}
    for key, val in request.POST.items():
        if not key.startswith("buffer_"):
            continue

        barcode = key[len("buffer_"):]
        value = val.strip()

        if value == "":
            buffer_override[barcode] = 0
            continue

        try:
            buffer_override[barcode] = float(value)
        except ValueError:
            pass

    try:
        result = run_plan(po_import_ids, buffer_override=buffer_override)
    except InactiveSkuOrderedError as e:
        return error_response(str(e), status=409)
    except ReconciliationError as e:
        return error_response(str(e), status=409)
    except DuplicateSubLocationError as e:
        return error_response(str(e), status=409)
    except Exception as e:
        return error_response(f"สร้างแผนล้มเหลว: {type(e).__name__}: {e}", status=500)

    if is_htmx:
        # ใช้ "replaceLocation" (custom event ที่ตัวเองทำ window.location.replace()) แทน HX-Redirect
        # ธรรมดา — HX-Redirect จะ "เพิ่ม" หน้าแผนใหม่เข้า browser history (window.location = url) ทำให้
        # หน้ากรอกยอดเผื่อ (ที่ควรเป็นแค่ "ขั้นตอนแวะผ่าน" ไม่ใช่ปลายทาง) ยังค้างอยู่ใน history —
        # พอกดปุ่มย้อนกลับจากหน้าแผน จะเด้งไปหน้ากรอกยอดเผื่อแทนที่จะเป็นหน้า PO ที่กดสร้างแผนมาจริงๆ
        # .replace() แทนที่ entry ปัจจุบัน (หน้ากรอกยอดเผื่อ) เลย ทำให้กด back ข้ามไปหน้า PO ตรงๆ
        response = HttpResponse(status=200)
        response["HX-Trigger"] = json.dumps({
            "replaceLocation": {"url": reverse("cpall:view_plan", args=[result["plan_run_id"]])}
        })
        return response
    return redirect("cpall:view_plan", plan_run_id=result["plan_run_id"])


def edit_buffer_form(request, plan_run_id):
    """หน้าแก้ยอดเผื่อของแผนที่สร้างไปแล้ว — pre-fill ด้วยยอดเผื่อปัจจุบัน "ของแผนนี้จริงๆ" (ไม่ใช่ค่า
    ล่าสุดที่เคยบันทึกไว้ทั่วระบบแบบหน้ากรอกยอดเผื่อตอนสร้างแผนใหม่)"""
    plan_run = get_object_or_404(PlanRun, id=plan_run_id)

    import openpyxl

    from customers.cpall.logic.excel_export import TEMPLATE_PATH as PP_TEMPLATE_PATH
    from customers.cpall.logic.excel_export import _find_sku_header_rows as _find_pp_sku_header_rows
    from customers.cpall.logic.excel_export import get_sheet_name as get_pp_sheet_name
    from customers.cpall.logic.plan_runner import get_current_buffer_by_barcode
    from customers.cpall.models import ProductMaster

    wb = openpyxl.load_workbook(PP_TEMPLATE_PATH)
    ws = wb[get_pp_sheet_name()]
    header_rows = _find_pp_sku_header_rows(ws)

    # เหตุผลเดียวกับ buffer_form() — ดึง SKU จาก ProductMaster (DB) ที่ active เป็นหลัก ไม่ใช่จาก
    # Template ตรงๆ (2025-09-05)
    active_barcodes = set(ProductMaster.objects.filter(is_active=True).values_list("barcode", flat=True))
    barcodes = [
        bc for bc, _ in sorted(header_rows.items(), key=lambda kv: kv[1])
        if bc in active_barcodes
    ]

    current_buffer = get_current_buffer_by_barcode(plan_run_id)
    name_lookup = {s.barcode: s.name_th for s in ProductMaster.objects.filter(barcode__in=barcodes)}
    sku_rows = [
        {"barcode": bc, "name_th": name_lookup.get(bc, bc), "current_buffer": current_buffer.get(bc, 0)}
        for bc in barcodes
    ]

    return render(request, "cpall/edit_buffer_form.html", {
        "plan_run_id": plan_run_id, "plan_name": plan_run.get_short_label(), "sku_rows": sku_rows,
    })


def edit_buffer_form_submit(request, plan_run_id):
    if request.method != "POST":
        return redirect("cpall:view_plan", plan_run_id=plan_run_id)

    is_htmx = request.headers.get("HX-Request") == "true"

    def error_response(message, status=400):
        if is_htmx:
            response = HttpResponse(status=status)
            response["HX-Trigger"] = json.dumps({"toast": {"message": message, "level": "error"}})
            return response
        return render(request, "cpall/plan_error.html", {"error": message})

    buffer_override = {}

    for key, val in request.POST.items():
        if not key.startswith("buffer_"):
            continue

        barcode = key[len("buffer_"):]
        value = val.strip()

        if value == "":
            buffer_override[barcode] = 0
            continue

        try:
            buffer_override[barcode] = float(value)
        except ValueError:
            pass

    try:
        edit_buffer_and_regenerate(plan_run_id, buffer_override)
    except PlanRun.DoesNotExist:
        return error_response("ไม่พบแผนนี้ (อาจถูกลบไปแล้ว)", status=404)
    except InactiveSkuOrderedError as e:
        return error_response(str(e), status=409)
    except Exception as e:
        return error_response(f"แก้ยอดเผื่อล้มเหลว: {type(e).__name__}: {e}", status=500)

    if is_htmx:
        response = HttpResponse(status=200)
        response["HX-Trigger"] = json.dumps({
            "toast": {"message": "อัปเดตยอดเผื่อและคำนวณแผนใหม่แล้ว", "level": "success"},
            "goBackAfterSave": {},
        })
        return response
    return redirect("cpall:view_plan", plan_run_id=plan_run_id)


def _snapshot_plan_results(plan_run_id):
    """
    เก็บ snapshot ยอดของแผนไว้เทียบก่อน/หลังเพิ่ม PO (Add PO feature — 2025-09-12)
    key = (sheet_type, group_name, barcode, column_label) -> qty
    ใช้ตรวจว่าการเพิ่ม PO ไปกระทบ "ช่องที่มีข้อมูลอยู่แล้ว" หรือไม่ — ตามหลักการที่ตกลงกันไว้ว่า
    การเพิ่ม PO ควรไปเติมเฉพาะช่องที่ยังว่าง ไม่ควรแก้ข้อมูลเดิม ถ้ากระทบต้องแจ้งให้คนทำเลือกเอง
    """
    from customers.cpall.models import PlanSkuResult

    return {
        (r.sheet_type, r.group_name or "", r.barcode, r.column_label): r.qty
        for r in PlanSkuResult.objects.filter(plan_run_id=plan_run_id, qty__isnull=False)
    }


def _diff_plan_results(before: dict, after: dict) -> list[dict]:
    """
    เทียบ snapshot ก่อน/หลัง — คืนเฉพาะ "ช่องเดิมที่เปลี่ยนไปหรือหายไป" เท่านั้น
    ช่องใหม่ที่เพิ่มเข้ามา (จาก PO ที่เพิ่งเพิ่ม) ไม่ถือว่ากระทบ เพราะเป็นจุดประสงค์ของการเพิ่ม PO อยู่แล้ว
    """
    impacts = []
    for key, old_qty in before.items():
        sheet_type, group_name, barcode, column_label = key
        if key not in after:
            impacts.append({
                "sheet_type": sheet_type, "group_name": group_name, "barcode": barcode,
                "column_label": column_label, "old_qty": old_qty, "new_qty": None, "kind": "หายไป",
            })
        elif after[key] != old_qty:
            impacts.append({
                "sheet_type": sheet_type, "group_name": group_name, "barcode": barcode,
                "column_label": column_label, "old_qty": old_qty, "new_qty": after[key],
                "kind": "เปลี่ยนค่า",
            })
    return impacts


def add_po_form(request, plan_run_id):
    """หน้าเลือก PO ที่จะเพิ่มเข้าแผนเดิม (Add PO feature) — เลือกจาก PO ที่ import ไว้แล้ว หรือกด
    นำเข้าไฟล์ใหม่ (ซึ่งจะพากลับมาที่แผนนี้เองหลัง import เสร็จ)"""
    from customers.cpall.logic.po_parser import list_po_imports

    plan_run = get_object_or_404(PlanRun, id=plan_run_id)
    current_po_ids = set(plan_run.po_imports.values_list("id", flat=True))
    available_pos = [po for po in list_po_imports(limit=100) if po["id"] not in current_po_ids]
    current_pos = [
        {"id": p.id, "display_filename": os.path.basename(p.source_filename),
         "production_date": p.production_date, "po_date": p.po_date}
        for p in plan_run.po_imports.all()
    ]
    return render(request, "cpall/add_po_form.html", {
        "plan_run_id": plan_run_id,
        "plan_name": plan_run.get_short_label(),
        "current_pos": current_pos,
        "available_pos": available_pos,
        "form": ImportPOForm(),
    })


def add_po_buffer(request, plan_run_id):
    """หน้ากรอกยอดเผื่อของ flow เพิ่ม PO — ใช้ template เดียวกับหน้าแก้ยอดเผื่อปกติทุกประการ
    (ค่าปัจจุบันของแผนนี้ pre-fill ไว้ให้แล้ว ถ้าไม่แก้อะไรก็กดคำนวณได้เลย) ต่างแค่พก new_po_ids
    ติดไปด้วยเป็น hidden field — ยังไม่ผูก PO เข้าแผนจริงจนกว่าจะกดคำนวณ"""
    import openpyxl

    from customers.cpall.logic.excel_export import TEMPLATE_PATH as PP_TEMPLATE_PATH
    from customers.cpall.logic.excel_export import _find_sku_header_rows as _find_pp_sku_header_rows
    from customers.cpall.logic.excel_export import get_sheet_name as get_pp_sheet_name
    from customers.cpall.logic.plan_runner import get_current_buffer_by_barcode
    from customers.cpall.models import ProductMaster

    plan_run = get_object_or_404(PlanRun, id=plan_run_id)
    new_po_ids_str = request.GET.get("new_po_ids", "")
    new_po_ids = [int(x) for x in new_po_ids_str.split(",") if x.strip()]
    if not new_po_ids:
        return redirect("cpall:add_po_form", plan_run_id=plan_run_id)

    wb = openpyxl.load_workbook(PP_TEMPLATE_PATH)
    ws = wb[get_pp_sheet_name()]
    header_rows = _find_pp_sku_header_rows(ws)
    active_barcodes = set(ProductMaster.objects.filter(is_active=True).values_list("barcode", flat=True))
    barcodes = [bc for bc, _ in sorted(header_rows.items(), key=lambda kv: kv[1]) if bc in active_barcodes]

    current_buffer = get_current_buffer_by_barcode(plan_run_id)
    name_lookup = {s.barcode: s.name_th for s in ProductMaster.objects.filter(barcode__in=barcodes)}
    sku_rows = [
        {"barcode": bc, "name_th": name_lookup.get(bc, bc), "current_buffer": current_buffer.get(bc, 0)}
        for bc in barcodes
    ]
    new_pos = [
        {"display_filename": os.path.basename(p.source_filename),
         "production_date": p.production_date, "po_date": p.po_date}
        for p in PoImport.objects.filter(id__in=new_po_ids)
    ]
    return render(request, "cpall/edit_buffer_form.html", {
        "plan_run_id": plan_run_id, "plan_name": plan_run.get_short_label(), "sku_rows": sku_rows,
        "add_po_mode": True, "new_po_ids_str": ",".join(str(i) for i in new_po_ids), "new_pos": new_pos,
    })


def add_po_submit(request, plan_run_id):
    """
    ผูก PO ใหม่เข้าแผนเดิม + คำนวณใหม่ทั้งชุด (Add PO feature — 2025-09-12)

    ทำใน transaction เดียวเสมอ แล้วเทียบ snapshot ก่อน/หลัง:
      - ถ้าไม่มีช่องเดิมเปลี่ยนเลย  -> commit ปกติ จบเลย (กรณีปกติของ workflow รอบเย็น->รอบเช้า)
      - ถ้ามีช่องเดิมเปลี่ยน/หายไป -> rollback ทั้งหมด แล้วโชว์ว่าเปลี่ยนอะไรบ้าง ให้คนทำเลือกเองว่า
        จะยืนยันทำต่อไหม (ส่ง force=1 กลับมา) — ตามหลักการที่ตกลงกันไว้ว่า "ไม่แก้ข้อมูลเดิมเงียบๆ"
    """
    from django.db import transaction

    from customers.cpall.logic.grouping import check_duplicate_sub_locations
    from customers.cpall.models import PlanRunImport

    if request.method != "POST":
        return redirect("cpall:view_plan", plan_run_id=plan_run_id)

    is_htmx = request.headers.get("HX-Request") == "true"
    plan_run = get_object_or_404(PlanRun, id=plan_run_id)

    def error_response(message, status=400):
        if is_htmx:
            response = HttpResponse(status=status)
            response["HX-Trigger"] = json.dumps({"toast": {"message": message, "level": "error"}})
            return response
        return render(request, "cpall/plan_error.html", {"error": message})

    new_po_ids = [int(x) for x in request.POST.get("new_po_ids_str", "").split(",") if x.strip()]
    if not new_po_ids:
        return error_response("ไม่พบ PO ที่จะเพิ่ม")

    existing_ids = set(plan_run.po_imports.values_list("id", flat=True))
    new_po_ids = [i for i in new_po_ids if i not in existing_ids]
    if not new_po_ids:
        return error_response("PO ที่เลือกอยู่ในแผนนี้อยู่แล้ว")

    # จุดส่งย่อยซ้ำกันระหว่างรอบ PO = มักเป็นการเผลอเพิ่มไฟล์ซ้ำ/ไฟล์ผิด — บล็อกเหมือน flow สร้างแผน
    combined_ids = list(existing_ids) + new_po_ids
    duplicates = check_duplicate_sub_locations(combined_ids)
    if duplicates:
        details = "; ".join(
            f"{d['sub_location']} (PO Import {', '.join(map(str, d['po_import_ids']))})" for d in duplicates
        )
        return error_response(f"เพิ่ม PO ไม่ได้ — จุดส่งย่อยซ้ำกันในหลายรอบ PO: {details}", status=409)

    buffer_override = {}
    for key, val in request.POST.items():
        if not key.startswith("buffer_"):
            continue
        barcode = key[len("buffer_"):]
        value = val.strip()
        if value == "":
            buffer_override[barcode] = 0
            continue
        try:
            buffer_override[barcode] = float(value)
        except ValueError:
            pass

    force = request.POST.get("force") == "1"
    before = _snapshot_plan_results(plan_run_id)

    class _ImpactDetected(Exception):
        """ใช้บังคับ rollback เท่านั้น ไม่ใช่ error จริง"""

        def __init__(self, impacts):
            self.impacts = impacts

    try:
        with transaction.atomic():
            for po_id in new_po_ids:
                PlanRunImport.objects.create(plan_run_id=plan_run_id, po_import_id=po_id)
            edit_buffer_and_regenerate(plan_run_id, buffer_override)

            if not force:
                impacts = _diff_plan_results(before, _snapshot_plan_results(plan_run_id))
                if impacts:
                    raise _ImpactDetected(impacts)
    except _ImpactDetected as e:
        # rollback แล้ว — แผนเดิมยังอยู่ครบเหมือนเดิมทุกประการ ให้คนทำตัดสินใจเอง
        return render(request, "cpall/add_po_impact.html", {
            "plan_run_id": plan_run_id, "plan_name": plan_run.get_short_label(),
            "impacts": e.impacts[:100], "impact_total": len(e.impacts),
            "new_po_ids_str": ",".join(str(i) for i in new_po_ids),
            "buffer_override": buffer_override,
        }, status=409)
    except InactiveSkuOrderedError as e:
        return error_response(str(e), status=409)
    except Exception as e:
        return error_response(f"เพิ่ม PO ไม่สำเร็จ: {type(e).__name__}: {e}", status=500)

    if is_htmx:
        response = HttpResponse(status=200)
        response["HX-Trigger"] = json.dumps({
            "toast": {"message": f"เพิ่ม PO {len(new_po_ids)} รอบ และคำนวณแผนใหม่แล้ว", "level": "success"},
            "replaceLocation": {"url": reverse("cpall:view_plan", args=[plan_run_id])},
        })
        return response
    return redirect("cpall:view_plan", plan_run_id=plan_run_id)


def plan_note_submit(request, plan_run_id):
    """บันทึกหมายเหตุของแผน (Add PO feature — 2025-09-12)"""
    if request.method != "POST":
        return redirect("cpall:view_plan", plan_run_id=plan_run_id)
    plan_run = get_object_or_404(PlanRun, id=plan_run_id)
    plan_run.note = request.POST.get("note", "").strip() or None
    plan_run.save(update_fields=["note"])
    if request.headers.get("HX-Request") == "true":
        response = HttpResponse(status=200)
        response["HX-Trigger"] = json.dumps({"toast": {"message": "บันทึกหมายเหตุแล้ว", "level": "success"}})
        return response
    return redirect("cpall:view_plan", plan_run_id=plan_run_id)


def _set_download_filename(response, filename):
    """
    ตั้งชื่อไฟล์ดาวน์โหลดแบบรองรับภาษาไทยให้ถูกต้องตามมาตรฐาน RFC 6266 (filename*=UTF-8''...)
    ไม่ใช้ resp["Content-Disposition"] = f'...filename="{filename}"...' ตรงๆ เพราะ Django จะ auto-encode
    ทั้ง header เป็น MIME encoded-word แบบเก่า (=?utf-8?b?...?=) ทันทีที่เจอตัวอักษรที่ไม่ใช่ ASCII —
    ซึ่งเป็นมาตรฐานสำหรับ email header ไม่ใช่ HTTP header บางเบราว์เซอร์/ตัวจัดการดาวน์โหลด parse ผิด
    แล้ว fallback ไปใช้ชื่อ "download" เฉยๆ (ตรงกับปัญหาที่เจอจริง) — เข้ารหัสเองแบบ RFC 6266 แทน พร้อม
    ชื่อสำรอง ASCII ควบคู่กันไปด้วย เผื่อเบราว์เซอร์เก่ามากๆ ไม่รู้จัก filename*= เลย (นามสกุลของชื่อ
    สำรองต้องตรงกับไฟล์จริงเสมอ — ไม่ hardcode .xlsx เพราะไฟล์นี้อาจเป็น .zip ก็ได้)
    """
    ext = os.path.splitext(filename)[1] or ".xlsx"
    encoded = quote(filename)
    response["Content-Disposition"] = f"attachment; filename=\"download{ext}\"; filename*=UTF-8''{encoded}"


def view_plan(request, plan_run_id):
    detail = get_plan_run_detail(plan_run_id)
    if detail is None:
        return render(request, "cpall/plan_not_found.html", {"plan_run_id": plan_run_id}, status=404)
    plan_name = PlanRun.objects.get(id=plan_run_id).get_short_label()
    skipped_skus = get_skipped_skus(plan_run_id)
    return render(request, "cpall/plan_view.html", {
        "plan": detail, "plan_name": plan_name, "skipped_skus": skipped_skus,
    })


def download_production(request, plan_run_id):
    detail = get_plan_run_detail(plan_run_id)
    if detail is None or not detail["production_plan_path"]:
        raise Http404
    filename = f"{PlanRun.objects.get(id=plan_run_id).get_display_name(prefix='แพลน')}.xlsx"

    # ลองสร้างไฟล์ใหม่จากข้อมูลดิบก่อนเสมอ (มีสูตรจริงครบ ตรงกับที่ตกลงกันไว้) — ถ้าทำไม่ได้ (แผนเก่า
    # ก่อนมีระบบ data-first) fallback ไปเสิร์ฟไฟล์ที่เก็บไว้บนดิสก์แบบเดิม (ยังไม่ได้ลบไฟล์เก่าทิ้งในเฟสนี้)
    try:
        content = regenerate_production_plan_bytes(plan_run_id)
        response = HttpResponse(
            content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        _set_download_filename(response, filename)
        return response
    except PlanRegenerateError:
        pass

    response = FileResponse(open(detail["production_plan_path"], "rb"),
                             content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _set_download_filename(response, filename)
    return response


def download_logistic(request, plan_run_id, group_name):
    detail = get_plan_run_detail(plan_run_id)
    if detail is None:
        raise Http404
    match = next((lp for lp in detail["logistic_plans"] if lp["group_name"] == group_name), None)
    if match is None:
        raise Http404  # กลุ่มนี้ไม่มีอยู่ในระบบเลย (ชื่อกลุ่มผิด/URL ปลอม) — 404 ถูกต้องแล้ว
    if match["status"] != "success" or not match["file_path"]:
        # เหตุผลเดียวกับ view_logistic_table — โชว์ error message จริงแทน 404 เปล่าที่งงว่าเกิดอะไรขึ้น
        reason = match.get("error_message") or ("ไม่มีข้อมูลของกลุ่มนี้ในรอบ PO ที่เลือก" if match["status"] == "skipped" else "ไม่ทราบสาเหตุ")
        return render(request, "cpall/plan_error.html",
                       {"error": f"ดาวน์โหลดกลุ่ม '{group_name}' ไม่สำเร็จ: {reason}"})
    filename = f"{PlanRun.objects.get(id=plan_run_id).get_display_name(prefix=group_name)}.xlsx"

    try:
        content = regenerate_logistic_plan_bytes(plan_run_id, group_name)
        response = HttpResponse(
            content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        _set_download_filename(response, filename)
        return response
    except PlanRegenerateError:
        pass

    response = FileResponse(open(match["file_path"], "rb"),
                             content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _set_download_filename(response, filename)
    return response


def view_po_detail(request, po_import_id):
    detail = get_po_detail(po_import_id)
    if detail is None:
        raise Http404
    return render(request, "cpall/po_detail.html", {"po": detail})


def download_po(request, po_import_id):
    """สร้างไฟล์ PO ใหม่จากข้อมูลใน database (ครบทุกคอลัมน์เหมือนต้นฉบับ ทดสอบยืนยันตรงกันทุกเซลล์)
    — ไม่ได้เก็บไฟล์ต้นฉบับไว้ถาวรอีกต่อไป แต่ถ้าเป็น PO เก่าที่ยังมีไฟล์ค้างอยู่จากก่อนเปลี่ยนมาใช้
    ระบบนี้ ให้เสิร์ฟไฟล์เดิมนั้นตรงๆ ก่อน (ของจริงย่อมดีกว่าของสร้างใหม่เสมอถ้ามีอยู่จริง)"""
    detail = get_po_detail(po_import_id)
    if detail is None:
        raise Http404

    source_path = detail["source_filename"]
    if source_path and os.path.exists(source_path):
        response = FileResponse(
            open(source_path, "rb"),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        _set_download_filename(response, detail["display_filename"])
        return response

    try:
        content = regenerate_po_file_bytes(po_import_id)
    except PORegenerateError:
        raise Http404
    response = HttpResponse(
        content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    _set_download_filename(response, detail["display_filename"])
    return response


def download_all_zip(request, plan_run_id):
    """โหลดทั้งแผน (Production Plan + Logistic Plan ทุกกลุ่มที่สำเร็จ) รวมเป็นไฟล์ ZIP เดียว —
    ใช้ตัวสร้างไฟล์เดียวกับการโหลดทีละไฟล์ทุกอย่าง (มีสูตรจริงครบเหมือนกัน) แค่รวมเข้า zip ในหน่วยความจำ
    ไม่ผ่านไฟล์ชั่วคราวบนดิสก์เลย (เหมือน pattern ที่ใช้กับการโหลดทีละไฟล์)"""
    detail = get_plan_run_detail(plan_run_id)
    if detail is None:
        raise Http404
    plan = PlanRun.objects.get(id=plan_run_id)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        if detail["production_plan_status"] == "success":
            try:
                content = regenerate_production_plan_bytes(plan_run_id)
            except PlanRegenerateError:
                content = None
                if detail["production_plan_path"] and os.path.exists(detail["production_plan_path"]):
                    with open(detail["production_plan_path"], "rb") as f:
                        content = f.read()
            if content:
                zf.writestr(f"{plan.get_display_name(prefix='แพลน')}.xlsx", content)

        for lp in detail["logistic_plans"]:
            if lp["status"] != "success":
                continue
            try:
                content = regenerate_logistic_plan_bytes(plan_run_id, lp["group_name"])
            except PlanRegenerateError:
                content = None
                if lp["file_path"] and os.path.exists(lp["file_path"]):
                    with open(lp["file_path"], "rb") as f:
                        content = f.read()
            if content:
                zf.writestr(f"{plan.get_display_name(prefix=lp['group_name'])}.xlsx", content)

        # แจ้งเตือนถ้ามีไฟล์ที่ขาดหายไปจาก ZIP นี้ — เดิมไม่มีการแจ้งเตือนเลย ZIP ดาวน์โหลดสำเร็จปกติ
        # (status 200) ทั้งที่ขาดไฟล์ไปเงียบๆ Admin ไม่มีทางรู้ตัวจนกว่าจะนับไฟล์เอง (เจอบั๊กนี้จริงจาก
        # การทดสอบ end-to-end — อันตรายกว่าเจอ error ตรงๆ เสียอีก เพราะไม่รู้ตัวว่าขาดอะไรไปบ้าง) —
        # ใส่เป็นไฟล์ .txt แยกต่างหากในซิป ไม่กระทบไฟล์ Excel อื่นเลย
        missing_notes = []
        if detail["production_plan_status"] != "success":
            missing_notes.append(f"แพลนผลิต — {detail.get('production_plan_error') or 'ไม่ทราบสาเหตุ'}")
        for lp in detail["logistic_plans"]:
            if lp["status"] == "failed":
                missing_notes.append(f"แพลนกระจาย {lp['group_name']} — {lp.get('error_message') or 'ไม่ทราบสาเหตุ'}")
        if missing_notes:
            note_text = (
                "ไฟล์ต่อไปนี้ไม่ได้รวมอยู่ใน ZIP นี้ (สร้างไม่สำเร็จ):\n\n"
                + "\n\n".join(f"- {n}" for n in missing_notes)
            )
            zf.writestr("⚠️ อ่านก่อน - มีไฟล์ที่ขาดหายไป.txt", note_text)

    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/zip")
    zip_filename = f"{plan.get_display_name(prefix='แพลนทั้งหมด')}.zip"
    _set_download_filename(response, zip_filename)
    return response


# ---------- จัดการ Template (ดาวน์โหลด/อัปโหลด/ดูประวัติเวอร์ชัน/กู้คืน/ลบ) ----------

def template_list(request):
    """
    Template Group tab (default) — Feature 1, 2025-09-10 — แสดงรายการชุด Template ทั้งหมด
    """
    from customers.cpall.logic.template_manager import ensure_bootstrap_group
    from customers.cpall.models import TemplateGroup

    ensure_bootstrap_group()
    groups = TemplateGroup.objects.prefetch_related("items__template_version").all()
    return render(request, "cpall/template_list.html", {"groups": groups, "active_tab": "group"})


def template_version_list(request):
    """Template Versions tab — เนื้อหาเดิมทั้งหมดจากก่อนมี Feature 1 (ย้ายมาจาก template_list เดิม)"""
    templates = list_templates()
    return render(request, "cpall/_template_version_list_page.html", {"templates": templates})


def template_group_detail(request, group_id):
    """หน้ารายละเอียดชุด Template — แสดงสมาชิกในกลุ่ม + ผล validate_group_consistency() ล่าสุด"""
    from customers.cpall.models import TemplateGroup

    group = get_object_or_404(
        TemplateGroup.objects.prefetch_related("items__template_version"), id=group_id,
    )

    # production_version = None
    # logistic_versions = []
    # for item in group.items.all():
    #     if item.template_version.template_key == "production_plan":
    #         production_version = item.template_version
    #     else:
    #         logistic_versions.append(item.template_version)

    production_version, logistic_versions = get_group_template_versions(group)

    registry = get_template_registry()

    for version in logistic_versions:
        version.display_label = registry[version.template_key]["label"]

    consistency = None
    if production_version:
        try:
            consistency = validate_group_consistency(production_version, logistic_versions)
        except Exception as e:
            consistency = {"is_consistent": False, "error": str(e)}

    return render(request, "cpall/template_group_detail.html", {
        "group": group,
        "production_version": production_version,
        "logistic_versions": logistic_versions,
        "consistency": consistency,
    })


def _template_group_form_context():
    """เตรียมข้อมูล template ทั้งหมด (key + เวอร์ชันที่มีอยู่) สำหรับหน้าสร้าง/แก้ไข Group — ใช้ร่วมกัน
    ทั้ง GET (แสดงฟอร์ม) — แยกจาก view เพื่อไม่ต้องเขียนซ้ำระหว่างหน้าสร้างกับหน้าแก้ไข"""
    from customers.cpall.models import TemplateVersion

    registry = get_template_registry()
    templates = []
    for key, info in registry.items():
        versions = list(TemplateVersion.objects.filter(template_key=key).order_by("-version_number"))
        templates.append({"key": key, "label": info["label"], "kind": info["kind"], "versions": versions})
    return templates


def template_group_create(request):
    """สร้างชุด Template ใหม่ — เลือกเวอร์ชันที่มีอยู่ หรืออัปโหลดไฟล์ใหม่ได้ต่อ template (Feature 1)
    บันทึกแค่จัดสมาชิกในกลุ่มเท่านั้น ไม่แตะ ProductMaster เลย (เช็คแยกตอน Activate — Step ถัดไป)"""
    if request.method == "GET":
        return render(request, "cpall/template_group_form.html", {
            "templates": _template_group_form_context(), "group": None,
        })
    return _template_group_save(request, group=None)


def template_group_edit(request, group_id):
    """แก้ไขชุด Template ที่มีอยู่ — ใช้หน้า/logic เดียวกับสร้างใหม่ (ตามที่ตกลงกันไว้ว่าให้แยกหน้า
    แก้ไขจากหน้าอื่นชัดเจน กัน Admin ทำงานสับสน) — ถ้าชุดนี้ active อยู่ แสดง banner เตือนในฟอร์ม"""
    from customers.cpall.models import TemplateGroup

    group = get_object_or_404(TemplateGroup.objects.prefetch_related("items__template_version"), id=group_id)
    if request.method == "GET":
        current_version_ids = {item.template_version_id for item in group.items.all()}
        return render(request, "cpall/template_group_form.html", {
            "templates": _template_group_form_context(), "group": group,
            "current_version_ids": current_version_ids,
        })
    return _template_group_save(request, group=group)


def _template_group_save(request, group):
    """บันทึกฟอร์มสร้าง/แก้ไข Group จริง — รองรับทั้ง 'เลือกเวอร์ชันที่มีอยู่' และ 'อัปโหลดไฟล์ใหม่'
    ต่อ template slot (production_plan + logistic_<group> ทุกกลุ่มที่ active) — atomic ทั้งหมด (ถ้า
    ไฟล์ไหนอัปโหลดไม่ผ่าน ไม่บันทึกอะไรเลย ไม่เหลือ Group ค้างครึ่งๆ กลางๆ)"""
    from django.db import transaction

    from customers.cpall.logic.template_manager import TemplateValidationError, upload_version_for_group
    from customers.cpall.models import TemplateGroup, TemplateGroupItem

    is_htmx = request.headers.get("HX-Request") == "true"

    def error_response(message, status=400):
        if is_htmx:
            response = HttpResponse(status=status)
            response["HX-Trigger"] = json.dumps({"toast": {"message": message, "level": "error"}})
            return response
        return render(request, "cpall/template_group_form.html", {
            "templates": _template_group_form_context(), "group": group, "error": message,
        })

    name = request.POST.get("name", "").strip()
    note = request.POST.get("note", "").strip()
    if not name:
        return error_response("กรุณากรอกชื่อชุด")

    registry = get_template_registry()
    temp_files_to_cleanup = []
    try:
        version_ids = []
        for key in registry:
            mode = request.POST.get(f"{key}_mode")
            if mode == "existing":
                version_id = request.POST.get(f"{key}_existing_version")
                if version_id:
                    version_ids.append(int(version_id))
            elif mode == "upload":
                uploaded_file = request.FILES.get(f"{key}_upload_file")
                if uploaded_file:
                    ext = os.path.splitext(uploaded_file.name)[1].lower()
                    if ext != ".xlsx":
                        return error_response(f"ไฟล์ของ '{registry[key]['label']}' ต้องเป็น .xlsx เท่านั้น")
                    os.makedirs(TEMP_UPLOAD_DIR, exist_ok=True)
                    temp_path = os.path.join(TEMP_UPLOAD_DIR, f"group_{key}_{uploaded_file.name}")
                    with open(temp_path, "wb") as f:
                        for chunk in uploaded_file.chunks():
                            f.write(chunk)
                    temp_files_to_cleanup.append(temp_path)
                    new_version = upload_version_for_group(key, temp_path, original_filename=uploaded_file.name)
                    temp_files_to_cleanup.remove(temp_path)  # upload_version_for_group ย้ายไฟล์ไปแล้ว
                    version_ids.append(new_version.id)

        with transaction.atomic():
            if group is None:
                from customers.cpall.logic.db import get_cpall_customer_id
                group = TemplateGroup.objects.create(customer_id=get_cpall_customer_id(), name=name, note=note)
            else:
                group.name = name
                group.note = note
                group.save(update_fields=["name", "note"])
                TemplateGroupItem.objects.filter(template_group=group).delete()

            for vid in version_ids:
                TemplateGroupItem.objects.create(template_group=group, template_version_id=vid)

    except TemplateValidationError as e:
        return error_response(f"ไฟล์มีปัญหา: {e}")
    finally:
        for path in temp_files_to_cleanup:
            if os.path.exists(path):
                os.remove(path)

    if is_htmx:
        response = HttpResponse(status=200)
        response["HX-Trigger"] = json.dumps({
            "toast": {"message": "บันทึกชุด Template สำเร็จ", "level": "success"},
            "replaceLocation": {"url": reverse("cpall:template_group_detail", args=[group.id])},
        })
        return response
    return redirect("cpall:template_group_detail", group_id=group.id)

def template_group_reconcile_submit(request, group_id):
    if request.method != "POST":
        return HttpResponse(status=405)

    from customers.cpall.models import TemplateGroup

    group = get_object_or_404(
        TemplateGroup.objects.prefetch_related("items__template_version"),
        id=group_id,
    )

    production_version, logistic_versions = get_group_template_versions(group)

    if production_version is None:
        return HttpResponse(
            "ชุดนี้ยังไม่มีแพลนผลิต",
            status=400,
        )

    # ตรวจไฟล์กับไฟล์ใหม่อีกครั้ง
    consistency = validate_group_consistency(
        production_version,
        logistic_versions,
    )

    if not consistency["is_consistent"]:
        return HttpResponse(
            "Template ในชุดนี้ไม่สอดคล้องกันแล้ว กรุณาตรวจสอบใหม่",
            status=400,
        )

    # ตรวจ ProductMaster ใหม่จากข้อมูลปัจจุบัน
    all_versions = [production_version] + logistic_versions

    reconciles = [
        reconcile_template_version(
            key=v.template_key,
            version_id=v.id,
        )
        for v in all_versions
    ]

    reconcile_data = build_group_reconcile_actions(
        all_versions,
        reconciles,
    )

    # ห้ามมี conflict
    if reconcile_data["conflicts"]:
        return HttpResponse(
            "พบรายการ ProductMaster ที่มี action ขัดแย้งกัน",
            status=400,
        )

    server_actions = reconcile_data["actions"]

    # ตรวจ action ที่ส่งมาจากหน้าเว็บกับ action ที่ server คำนวณไว้
    submitted_actions = list(
        zip(
            request.POST.getlist("action"),
            request.POST.getlist("barcode"),
        )
    )

    submitted_action_keys = {
        (str(action).strip(), str(barcode).strip())
        for action, barcode in submitted_actions
    }

    server_action_keys = {
        (
            str(item["action"]).strip(),
            str(item["barcode"]).strip(),
        )
        for item in server_actions
    }

    if submitted_action_keys != server_action_keys:
        return HttpResponse(
            "รายการ ProductMaster เปลี่ยนแปลงแล้ว กรุณาตรวจสอบชุด Template ใหม่",
            status=400,
        )

    # ตรวจข้อมูลสำหรับ CREATE_PRODUCT
    create_data = {}

    for item in server_actions:
        if item["action"] != "CREATE_PRODUCT":
            continue

        barcode = str(item["barcode"]).strip()

        name_th = str(
            request.POST.get(f"name_th_{barcode}", "")
        ).strip()

        name_en = str(
            request.POST.get(f"name_en_{barcode}", "")
        ).strip()

        pack_size = str(
            request.POST.get(f"pack_size_{barcode}", "")
        ).strip()

        unit_price = str(
            request.POST.get(f"unit_price_{barcode}", "")
        ).strip()

        if not name_th:
            return HttpResponse(
                f"กรุณาระบุชื่อสินค้า TH สำหรับ Barcode {barcode}",
                status=400,
            )

        create_data[barcode] = {
            "name_th": name_th,
            "name_en": name_en,
            "pack_size": pack_size,
            "unit_price": unit_price,
        }

    # เตรียม ProductMaster actions ทั้งหมด
    product_actions = []

    for item in server_actions:
        action = item["action"]
        barcode = str(item["barcode"]).strip()

        action_data = {
            "action": action,
            "barcode": barcode,
        }

        if action == "CREATE_PRODUCT":
            data = create_data[barcode]

            action_data.update({
                "name_th": data["name_th"],
                "name_en": data["name_en"],
                "pack_size": data["pack_size"],
                "unit_price": data["unit_price"],
            })

        product_actions.append(action_data)

    # Commit ทุกอย่าง
    try:
        activate_template_group(
            group,
            product_actions=product_actions,
        )
    except Exception:
        return HttpResponse(
            "ไม่สามารถใช้ชุด Template ได้ กรุณาติดต่อผู้ดูแลระบบ",
            status=400,
        )

    return redirect(
        "cpall:template_group_detail",
        group_id=group.id,
    )


def template_group_activate(request, group_id):
    """
    ปุ่ม "ใช้ชุดนี้" — Activate ทั้ง Group พร้อมกัน

    ลำดับ:
    [1] validate_group_consistency()
        ไฟล์ในกลุ่มต้องสอดคล้องกันก่อน
        ไม่ผ่าน → block และให้ดูรายละเอียด

    [2] reconcile_template_version()
        ตรวจ ProductMaster ของทุก Template

    [3a] ไม่มี mismatch
        → activate_template_group() ทันที
        → กลับหน้า Group Detail

    [3b] มี mismatch
        → หน้า template_group_reconcile.html
        → Admin กรอกข้อมูล CREATE_PRODUCT
        → กด "ยืนยันใช้ชุดนี้"
        → submit ตรวจสอบใหม่และ activate
    """
    from customers.cpall.models import TemplateGroup

    if request.method != "POST":
        return HttpResponse(status=405)

    group = get_object_or_404(TemplateGroup.objects.prefetch_related("items__template_version"), id=group_id)
    # is_htmx = request.headers.get("HX-Request") == "true"

    # production_version = None
    # logistic_versions = []
    # for item in group.items.all():
    #     if item.template_version.template_key == "production_plan":
    #         production_version = item.template_version
    #     else:
    #         logistic_versions.append(item.template_version)

    production_version, logistic_versions = get_group_template_versions(group)

    if production_version is None:
        response = HttpResponse(status=400)
        response["HX-Trigger"] = json.dumps({
            "toast": {"message": "ชุดนี้ยังไม่มีแพลนผลิต — ใช้งานไม่ได้", "level": "error"},
        })
        return response

    # [1] ตรวจไฟล์กับไฟล์ก่อน — ไม่ผ่าน → popup error (ไม่ list รหัสสินค้ายาวๆ ในข้อความตรงๆ ตามที่
    # ขอ 2025-09-11 — สรุปจำนวนสั้นๆ แล้วให้กด "ดูรายละเอียด ↗" ไปหน้า Group detail แทน ที่มีรายละเอียด
    # ครบอยู่แล้ว (consistency คำนวณสดใหม่ทุกครั้งที่โหลดหน้า) — reuse pattern popup/detail_url เดิม
    # จาก base.html's alert modal (เคยใช้กับ missing_template_items ตอน import PO มาก่อนแล้ว)
    consistency = validate_group_consistency(production_version, logistic_versions)
    if not consistency["is_consistent"]:
        mismatch_count = (
            len(consistency["product_missing_in_logistic"]) + len(consistency["product_extra_in_logistic"])
            + len(consistency["sub_location_missing_in_logistic"]) + len(consistency["sub_location_extra_in_logistic"])
        )
        response = HttpResponse(status=400)
        response["HX-Trigger"] = json.dumps({
            "toast": {
                "message": (
                    f"ไฟล์ในชุด '{group.name}' ยังไม่ตรงกัน "
                    f"({mismatch_count} รายการ) "
                    "แก้ไขให้ตรงกันก่อนถึงจะใช้ชุดนี้ได้"
                ),
                "level": "error",
                "detail_url": reverse(
                    "cpall:template_group_consistency_detail",
                    args=[group.id],
                ),
                "detail_label": "ดูรายการที่ไม่ตรงกัน",
            },
        })
        return response

    # [2] ตรวจ ProductMaster ทีละ template ในกลุ่ม (ของเดิม)
    all_versions = [production_version] + logistic_versions
    reconciles = [
        reconcile_template_version(
            key=v.template_key,
            version_id=v.id,
        )
        for v in all_versions
    ]

    has_mismatch = any(
        reconcile["mismatch_count"] > 0
        for reconcile in reconciles
    )

    if has_mismatch:
        reconcile_data = build_group_reconcile_actions(
            all_versions,
            reconciles,
        )

        if reconcile_data["conflicts"]:
            return render(
                request,
                "cpall/template_group_reconcile.html",
                {
                    "group": group,
                    "actions": [],
                    "conflicts": reconcile_data["conflicts"],
                },
            )

        return render(
            request,
            "cpall/template_group_reconcile.html",
            {
                "group": group,
                "actions": reconcile_data["actions"],
                "conflicts": [],
            },
        )

    # [3a] ไม่มี mismatch เลย → activate ทันที
    try:
        activate_template_group(
            group,
            product_actions=[],
        )
    except Exception:
        response = HttpResponse(status=400)
        response["HX-Trigger"] = json.dumps({
            "toast": {
                "message": "ไม่สามารถใช้ชุด Template ได้ กรุณาติดต่อผู้ดูแลระบบ",
                "level": "error",
            },
        })
        return response

    return redirect(
        "cpall:template_group_detail",
        group_id=group.id,
    )

def template_download(request, key):
    registry = get_template_registry()
    if key not in registry:
        raise Http404
    path = registry[key]["path"]
    if not os.path.exists(path):
        raise Http404
    return FileResponse(open(path, "rb"), as_attachment=True, filename=os.path.basename(path))


def template_upload(request, key):
    if request.method != "POST":
        return redirect("cpall:template_list")
    registry = get_template_registry()
    if key not in registry:
        raise Http404

    is_htmx = request.headers.get("HX-Request") == "true"

    def error_response(message):
        if is_htmx:
            response = HttpResponse(status=400)
            response["HX-Trigger"] = json.dumps({"toast": {"message": message, "level": "error"}})
            return response
        return render(request, "cpall/template_upload_result.html",
                       {"success": False, "error": message, "key": key, "label": registry[key]["label"]})

    form = TemplateUploadForm(request.POST, request.FILES)
    if not form.is_valid():
        error_text = "; ".join(e for errs in form.errors.values() for e in errs)
        return error_response(error_text or "ยังไม่ได้เลือกไฟล์")
    new_file = form.cleaned_data["template_file"]

    os.makedirs(TEMP_UPLOAD_DIR, exist_ok=True)
    temp_path = os.path.join(TEMP_UPLOAD_DIR, f"{key}_{new_file.name}")
    with open(temp_path, "wb") as f:
        for chunk in new_file.chunks():
            f.write(chunk)

    try:
        result = upload_new_version(
            key,
            temp_path,
            original_filename=new_file.name,
        )
    except TemplateValidationError as e:
        # validate ไม่ผ่าน -> เวอร์ชัน/ไฟล์ live เดิมไม่ถูกแตะเลย ลบไฟล์ที่อัปโหลดมาทิ้ง
        if os.path.exists(temp_path):
            os.remove(temp_path)
        return error_response(str(e))

    # สำเร็จ -> ไปหน้าประวัติเวอร์ชัน
    if is_htmx:
        url = reverse("cpall:template_versions", args=[key])
        url += f"?new_version={result['version'].id}"

        response = HttpResponse(status=200)
        response["HX-Redirect"] = url
        return response

    return redirect(
        "cpall:template_versions",
        key=key,
    )


def template_versions(request, key):
    """หน้าประวัติทุกเวอร์ชันของ template นี้ — กู้คืน/ลบได้จากหน้านี้"""
    registry = get_template_registry()
    if key not in registry:
        raise Http404
    versions = list_versions(key)

    new_version_id = request.GET.get("new_version")

    try:
        new_version_id = int(new_version_id) if new_version_id else None
    except ValueError:
        new_version_id = None

    return render(request, "cpall/template_versions.html", {
        "key": key,
        "label": registry[key]["label"],
        "versions": versions,
        "new_version_id": new_version_id,
    })


# def template_version_restore(request, key, version_id):
#     if request.method != "POST":
#         return redirect("cpall:template_versions", key=key)
#     if key not in get_template_registry():
#         raise Http404
#     is_htmx = request.headers.get("HX-Request") == "true"
#     try:
#         # version = restore_to_version(key, version_id)
#         # toast = {"message": f"ใช้เวอร์ชัน {version.version_number} แล้ว", "level": "success"}
#         result = restore_to_version(key, version_id)
#         version = result["version"]

#         if result["status"] == "pending":
#             if is_htmx:
#                 response = render(
#                     request,
#                     "cpall/_template_version_reconcile.html",
#                     {
#                         "key": key,
#                         "version_id": version_id,
#                         "result": result["reconcile"],
#                     },
#                 )

#                 response["HX-Trigger"] = json.dumps({
#                     "toast": {
#                         "message": (
#                             f"เวอร์ชัน {version.version_number} "
#                             "ต้องตรวจสอบ ProductMaster ก่อนใช้งาน"
#                         ),
#                         "level": "warning",
#                     }
#                 })

#                 return response

#             return redirect("cpall:template_versions", key=key)
#         else:
#             toast = {
#                 "message": f"ใช้เวอร์ชัน {version.version_number} แล้ว",
#                 "level": "success",
#             }
#     except TemplateValidationError as e:
#         if not is_htmx:
#             return redirect("cpall:template_versions", key=key)
#         toast = {"message": str(e), "level": "error"}

#     if is_htmx:
#         versions = list_versions(key)
#         response = render(request, "cpall/_template_version_list.html", {"key": key, "versions": versions})
#         response["HX-Trigger"] = json.dumps({"toast": toast})
#         return response
#     return redirect("cpall:template_versions", key=key)

def template_version_restore(request, key, version_id):
    if request.method != "POST":
        return redirect("cpall:template_versions", key=key)

    if key not in get_template_registry():
        raise Http404

    is_htmx = request.headers.get("HX-Request") == "true"

    try:
        result = restore_to_version(key, version_id)
        version = result["version"]

        if result["status"] == "pending":
            if is_htmx:
                response = render(
                    request,
                    "cpall/_template_version_reconcile.html",
                    {
                        "key": key,
                        "version_id": version_id,
                        "result": result["reconcile"],
                    },
                )

                response["HX-Trigger"] = json.dumps({
                    "toast": {
                        "message": (
                            f"เวอร์ชัน {version.version_number} "
                            "ต้องตรวจสอบ ProductMaster ก่อนใช้งาน"
                        ),
                        "level": "warning",
                    }
                })

                return response

            return redirect("cpall:template_versions", key=key)

        # ProductMaster ตรงกันแล้ว
        # ยังไม่ Active / Sync จนกว่า Admin จะยืนยัน
        if is_htmx:
            response = render(
                request,
                "cpall/_template_version_confirm.html",
                {
                    "key": key,
                    "version_id": version_id,
                    "version": version,
                },
            )

            return response

        return redirect("cpall:template_versions", key=key)

    except TemplateValidationError as e:
        if is_htmx:
            response = HttpResponse(status=400)
            response["HX-Trigger"] = json.dumps({
                "toast": {
                    "message": str(e),
                    "level": "error",
                }
            })
            return response

        return redirect("cpall:template_versions", key=key)

# def template_version_activate(request, key, version_id):
#     if request.method != "POST":
#         return HttpResponse(status=405)

#     if key not in get_template_registry():
#         raise Http404

#     is_htmx = request.headers.get("HX-Request") == "true"

#     try:
#         from customers.cpall.models import TemplateVersion

#         version = TemplateVersion.objects.get(
#             id=version_id,
#             template_key=key,
#         )

#         # ตรวจ ProductMaster ซ้ำอีกครั้งก่อนเปลี่ยนจริง
#         reconcile = reconcile_template_version(
#             key=key,
#             version_id=version_id,
#         )

#         if reconcile["mismatch_count"] > 0:
#             return HttpResponse(
#                 "ProductMaster ยังไม่ตรงกับ Template",
#                 status=400,
#             )

#         TemplateVersion.objects.filter(
#             template_key=key,
#             is_active=True,
#         ).update(is_active=False)

#         version.is_active = True
#         version.save(update_fields=["is_active"])

#         _sync_live_file(key, version)

#         if is_htmx:
#             versions = list_versions(key)

#             response = render(
#                 request,
#                 "cpall/_template_version_list.html",
#                 {
#                     "key": key,
#                     "versions": versions,
#                 },
#             )

#             response["HX-Trigger"] = json.dumps({
#                 "toast": {
#                     "message": (
#                         f"เปลี่ยนเป็นเวอร์ชัน "
#                         f"{version.version_number} แล้ว"
#                     ),
#                     "level": "success",
#                 }
#             })

#             return response

#         return redirect("cpall:template_versions", key=key)

#     except TemplateVersion.DoesNotExist:
#         raise Http404
#     except TemplateValidationError as e:
#         if is_htmx:
#             return HttpResponse(str(e), status=400)

#         return redirect("cpall:template_versions", key=key)

def template_version_delete(request, key, version_id):
    if request.method != "POST":
        return redirect("cpall:template_versions", key=key)
    registry = get_template_registry()
    if key not in registry:
        raise Http404
    is_htmx = request.headers.get("HX-Request") == "true"
    try:
        delete_version(key, version_id)
        toast = {"message": "ลบเวอร์ชันสำเร็จ", "level": "success"}
    except (TemplateInUseError, TemplateValidationError) as e:
        toast = {"message": str(e), "level": "error"}
        if not is_htmx:
            versions = list_versions(key)
            return render(request, "cpall/template_versions.html", {
                "key": key, "label": registry[key]["label"], "versions": versions, "error": str(e),
            })

    versions = list_versions(key)
    if is_htmx:
        response = render(request, "cpall/_template_version_list.html", {"key": key, "versions": versions})
        response["HX-Trigger"] = json.dumps({"toast": toast})
        return response
    return render(request, "cpall/template_versions.html", {
        "key": key, "label": registry[key]["label"], "versions": versions,
    })

def template_version_reconcile(request, key, version_id):
    if request.method != "GET":
        return HttpResponse(status=405)

    if key not in get_template_registry():
        raise Http404

    try:
        result = reconcile_template_version(
            key=key,
            version_id=version_id,
        )
    except TemplateValidationError as e:
        return HttpResponse(str(e), status=400)

    return render(
        request,
        "cpall/_template_version_reconcile.html",
        {
            "key": key,
            "version_id": version_id,
            "result": result,
        },
    )

def template_version_product_action(request, key, version_id):
    if request.method != "POST":
        return HttpResponse(status=405)

    if key not in get_template_registry():
        raise Http404

    is_htmx = request.headers.get("HX-Request") == "true"

    action = request.POST.get("action", "").strip()
    barcode = request.POST.get("barcode", "").strip()

    name_th = request.POST.get("name_th", "").strip()
    name_en = request.POST.get("name_en", "").strip()
    pack_size_str = request.POST.get("pack_size", "").strip()
    unit_price_str = request.POST.get("unit_price", "").strip()

    if action == "CREATE_PRODUCT":
        if not name_th:
            return HttpResponse(
                "กรุณากรอกชื่อสินค้า (ไทย)",
                status=400,
            )

        if not pack_size_str:
            return HttpResponse(
                "กรุณากรอก Pack Size",
                status=400,
            )

        if not unit_price_str:
            return HttpResponse(
                "กรุณากรอกราคา",
                status=400,
            )

        try:
            pack_size = int(pack_size_str)
            unit_price = float(unit_price_str)

            if pack_size <= 0 or unit_price < 0:
                raise ValueError

        except ValueError:
            return HttpResponse(
                "Pack Size หรือราคาไม่ถูกต้อง",
                status=400,
            )
    else:
        pack_size = None
        unit_price = None

    try:
        result = apply_product_master_action(
            action=action,
            barcode=barcode,
            name_th=name_th or None,
            name_en=name_en or None,
            pack_size=pack_size,
            unit_price=unit_price,
        )

        # ตรวจสอบ Template กับ ProductMaster ใหม่ทันที
        reconcile = reconcile_template_version(
            key=key,
            version_id=version_id,
        )

        if reconcile["mismatch_count"] == 0:
            from customers.cpall.models import TemplateVersion

            version = TemplateVersion.objects.get(
                id=version_id,
                template_key=key,
            )

            response = render(
                request,
                "cpall/_template_version_confirm.html",
                {
                    "key": key,
                    "version_id": version_id,
                    "version": version,
                },
            )
        else:
            response = render(
                request,
                "cpall/_template_version_reconcile.html",
                {
                    "key": key,
                    "version_id": version_id,
                    "result": reconcile,
                },
            )

        action_messages = {
            "CREATE_PRODUCT": "สร้างสินค้า",
            "ACTIVATE": "เปิดใช้งานสินค้า",
            "DEACTIVATE": "ปิดใช้งานสินค้า",
        }

        action_label = action_messages.get(
            result["action"],
            "ดำเนินการ",
        )

        # product_name = name_th or barcode
        product_name = result["product_name"]

        response["HX-Trigger"] = json.dumps({
            "toast": {
                "message": (
                    f"{action_label} “{product_name}” "
                    "ใน ProductMaster แล้ว"
                ),
                "level": "success",
            }
        })

        return response

    except TemplateValidationError as e:
        if is_htmx:
            response = HttpResponse(status=400)
            response["HX-Trigger"] = json.dumps({
                "toast": {
                    "message": str(e),
                    "level": "error",
                }
            })
            return response

        return redirect("cpall:template_versions", key=key)


# ---------- ดูตารางตัวเลขจริงในหน้าเว็บ (ไม่ต้องดาวน์โหลด Excel) ----------

def view_production_table(request, plan_run_id):
    detail = get_plan_run_detail(plan_run_id)
    if detail is None or detail["production_plan_status"] != "success":
        raise Http404
    # ลองอ่านจาก plan_sku_result ก่อน (ค่าจริงจาก LibreOffice) — ถ้าไม่มี (แผนเก่าก่อนมีระบบนี้ หรือ
    # extraction ตอนสร้างแผนล้มเหลว) ค่อย fallback ไปเปิดไฟล์ Excel + สูตรจำลอง Python แบบเดิม
    table = get_production_plan_table_from_db(plan_run_id)
    if not table["rows"]:
        table = get_production_plan_table(detail["production_plan_path"])
    # Django template ดึงค่าจาก dict ด้วย key ที่เป็นตัวแปรไม่ได้ตรงๆ -> จัดเป็น list ของ (qty, pack_text)
    # คู่กันตามลำดับคอลัมน์ไว้ก่อน เพื่อวนลูปแสดงคู่กันในเทมเพลตได้ง่ายๆ
    for row in table["rows"]:
        row["location_items"] = [
            {"qty": row["qty_by_location"].get(loc), "pack_text": row["pack_text_by_location"].get(loc)}
            for loc in table["sub_locations"]
        ]
    return render(request, "cpall/table_production.html", {"plan": detail, "table": table})


def view_logistic_table(request, plan_run_id, group_name):
    detail = get_plan_run_detail(plan_run_id)
    if detail is None:
        raise Http404
    lp = next((lp for lp in detail["logistic_plans"] if lp["group_name"] == group_name), None)
    if lp is None:
        raise Http404  # กลุ่มนี้ไม่มีอยู่ในระบบเลย (ชื่อกลุ่มผิด/URL ปลอม) — 404 ถูกต้องแล้ว
    if lp["status"] == "skipped":
        return render(request, "cpall/plan_error.html",
                       {"error": f"รอบ PO ที่เลือกไว้ไม่มีข้อมูลของกลุ่ม '{group_name}' เลย"})
    if lp["status"] == "failed":
        # เดิมตรงนี้ raise Http404 เฉยๆ ทำให้ Admin เจอหน้า 404 เปล่าไม่รู้สาเหตุ ทั้งที่ error_message
        # มีรายละเอียดเก็บไว้อยู่แล้วครบถ้วน (เช่น "เทมเพลตคอลัมน์ PO ไม่พอสำหรับรอบนี้") — เจอปัญหานี้
        # จริงจากการทดสอบ end-to-end แบบเต็มวงจร (กลุ่มหนึ่งสร้างไม่สำเร็จ แต่กลุ่มอื่นสำเร็จ แล้วกด
        # ดูตารางของกลุ่มที่ไม่สำเร็จ) — แก้ให้โชว์ error message จริงแทน
        return render(request, "cpall/plan_error.html",
                       {"error": f"สร้างตารางกลุ่ม '{group_name}' ไม่สำเร็จ: {lp.get('error_message') or 'ไม่ทราบสาเหตุ'}"})
    match = lp
    table = get_logistic_plan_table_from_db(plan_run_id, group_name)
    if not table["rows"]:
        table = get_logistic_plan_table(match["file_path"], group_name)
    for row in table["rows"]:
        row["location_items"] = [
            {"qty": row["qty_by_column"].get(col), "pack_text": row["pack_text_by_column"].get(col)}
            for col in table["columns"]
        ]
    total_baskets = sum(row["basket_total"] or 0 for row in table["rows"])
    basket_total_by_column = table.get("basket_total_by_column", {})  # ไม่มีถ้าเป็น fallback
    # ไปใช้ get_logistic_plan_table() (แผนเก่ามากที่ไม่มีข้อมูลใน plan_sku_result เลย) — ปล่อยว่างไว้
    # ก็พอ (แสดงแค่ "รวมตะกร้าทั้งหมด" แบบเดิม ไม่มีตะกร้าต่อคอลัมน์ให้แผนเก่ากลุ่มนี้)

    # เลข PO จริงเบื้องหลัง "PO1"/"PO2" แต่ละคอลัมน์ — เอาไว้แสดง tooltip เฉยๆ ถ้าหาไม่ได้ (เช่น PO
    # ต้นทางถูกลบไปแล้ว) ไม่ให้กระทบหน้าตารางเลย แค่ไม่มี tooltip ให้
    from customers.cpall.logic.logistic_plan_export import get_po_number_by_column_label

    po_import_ids = [po["id"] for po in detail["po_imports"]]
    try:
        po_number_by_column = get_po_number_by_column_label(po_import_ids, group_name)
    except Exception:
        po_number_by_column = {}
    # เตรียมคู่ (คอลัมน์, เลข PO) ไว้ล่วงหน้า — Django template lookup แบบ dict.{{ loop_var }} หา
    # key ชื่อ "loop_var" ตรงๆ ไม่ resolve ค่าตัวแปรให้ ต้องจับคู่มาก่อนแบบนี้แทน
    columns_with_po = [(col, po_number_by_column.get(col)) for col in table["columns"]]
    # เหตุผลเดียวกัน — เตรียมคู่ (คอลัมน์, ตะกร้าต่อคอลัมน์) ไว้ล่วงหน้าด้วย
    columns_with_basket = [(col, basket_total_by_column.get(col)) for col in table["columns"]]

    return render(request, "cpall/table_logistic.html", {
        "plan": detail, "table": table, "group_name": group_name, "total_baskets": total_baskets,
        "columns_with_po": columns_with_po, "columns_with_basket": columns_with_basket,
    })


# ---------- ลบ PO / ลบแผน ----------

def delete_po_import_view(request, po_import_id):
    if request.method != "POST":
        return redirect("cpall:po_list")
    is_htmx = request.headers.get("HX-Request") == "true"
    try:
        delete_po_import(po_import_id)
    except POInUseError as e:
        if is_htmx:
            # 409 = ไม่สำเร็จ -> HTMX จะไม่เอา response ไปแทนที่แถวเดิม (แถวยังอยู่ครบ) แค่โชว์ toast แดง
            response = HttpResponse(status=409)
            response["HX-Trigger"] = json.dumps({"toast": {"message": str(e), "level": "error"}})
            return response
        return render(request, "cpall/plan_error.html", {"error": str(e)})

    if is_htmx:
        # ตัวเปล่า status 200 -> hx-swap="outerHTML" เอาไปแทนที่ <tr> เดิม = แถวหายไปจากตารางทันที
        response = HttpResponse(status=200)
        response["HX-Trigger"] = json.dumps({"toast": {"message": "ลบ PO สำเร็จ", "level": "success"}})
        return response
    return redirect("cpall:po_list")


def delete_plan_run_view(request, plan_run_id):
    if request.method != "POST":
        return redirect("cpall:plan_list")
    delete_plan_run(plan_run_id)
    if request.headers.get("HX-Request") == "true":
        response = HttpResponse(status=200)
        response["HX-Trigger"] = json.dumps({"toast": {"message": "ลบแผนสำเร็จ", "level": "success"}})
        return response
    return redirect("cpall:plan_list")


def template_view(request, key):
    if key not in get_template_registry():
        raise Http404
    sheet_name = request.GET.get("sheet")
    try:
        grid = get_template_grid(key, sheet_name=sheet_name)
    except TemplateValidationError as e:
        raise Http404(str(e))
    return render(
        request,
        "cpall/template_view.html",
        {
            "key": key,
            "grid": grid,
            "back_url": reverse("cpall:template_version_list"),
            "back_label": "กลับ Template Versions",
        },
    )

def template_version_view(request, group_id, key, version_id):
    from customers.cpall.models import TemplateGroup, TemplateVersion

    if key not in get_template_registry():
        raise Http404

    group = get_object_or_404(
        TemplateGroup,
        id=group_id,
    )

    version = get_object_or_404(
        TemplateVersion,
        id=version_id,
        template_key=key,
    )

    if not group.items.filter(template_version=version).exists():
        raise Http404

    sheet_name = request.GET.get("sheet")

    try:
        grid = get_template_grid(
            key,
            sheet_name=sheet_name,
            filepath=version.file_path,
            label=version.original_filename or f"{key} v{version.version_number}",
        )
    except TemplateValidationError as e:
        raise Http404(str(e))

    return render(
        request,
        "cpall/template_view.html",
        {
            "key": key,
            "grid": grid,
            "version": version,
            "version_view": True,
            "group": group,
            "back_url": reverse(
                "cpall:template_group_detail",
                args=[group.id],
            ),
            "back_label": "กลับ Group Detail",
        },
    )


def template_version_download(request, key, version_id):
    from customers.cpall.models import TemplateVersion

    if key not in get_template_registry():
        raise Http404

    version = get_object_or_404(
        TemplateVersion,
        id=version_id,
        template_key=key,
    )

    if not version.file_path or not os.path.exists(version.file_path):
        raise Http404

    filename = (
        version.original_filename
        or f"{key}_v{version.version_number}.xlsx"
    )

    response = FileResponse(
        open(version.file_path, "rb"),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    _set_download_filename(response, filename)
    return response

def template_group_consistency_detail(request, group_id):
    from customers.cpall.models import TemplateGroup

    group = get_object_or_404(
        TemplateGroup.objects.prefetch_related("items__template_version"),
        id=group_id,
    )

    production_version, logistic_versions = get_group_template_versions(group)

    if production_version is None:
        return HttpResponse("ยังไม่มี Production Template", status=400)

    consistency = validate_group_consistency(
        production_version,
        logistic_versions,
    )

    return render(
        request,
        "cpall/template_group_consistency_detail.html",
        {
            "group": group,
            "consistency": consistency,
        },
    )

def template_version_history_view(request, key, version_id):
    from customers.cpall.models import TemplateVersion

    if key not in get_template_registry():
        raise Http404

    version = get_object_or_404(
        TemplateVersion,
        id=version_id,
        template_key=key,
    )

    sheet_name = request.GET.get("sheet")

    try:
        grid = get_template_grid(
            key,
            sheet_name=sheet_name,
            filepath=version.file_path,
            label=version.original_filename or f"{key} v{version.version_number}",
        )
    except TemplateValidationError as e:
        raise Http404(str(e))

    return render(
        request,
        "cpall/template_view.html",
        {
            "key": key,
            "grid": grid,
            "version": version,
            "version_view": True,
            "back_url": reverse("cpall:template_versions", args=[key]),
            "back_label": "กลับ Template Versions",
        },
    )